"""Training tab — dataset split, feature extraction, head training, fine-tuning, evaluation, TFLite export."""
import json
import threading
from pathlib import Path

# Module-level lock — persists across Streamlit reruns; prevents duplicate training threads.
_TRAIN_LOCK = threading.Lock()

import pandas as pd
import plotly.graph_objects as go
import streamlit as st
from plotly.subplots import make_subplots

from components.model_cache import load_cached_model
from components.ui_helpers import html, model_picker, detect_dir
from config import DISEASE_DISPLAY
from pipeline import (
    split_dataset, extract_features, train_head, fine_tune_model,
    evaluate_model, compute_tsne, check_class_balance, convert_model,
)
from streamlit_callback import StreamlitTrainCallback
from src.training.data import next_training_output_dir as _next_training_output_dir


# ---------------------------------------------------------------------------
# Colour palette shared by t-SNE and model-comparison charts
# ---------------------------------------------------------------------------

_PALETTE = [
    "#E63946", "#F4A261", "#2EC4B6", "#457B9D", "#A8DADC",
    "#6A4C93", "#FB8500", "#38B000", "#D62839", "#3A86FF",
]


# ---------------------------------------------------------------------------
# Learning-curve chart
# ---------------------------------------------------------------------------

def plot_learning_curves(history: dict) -> None:
    if not history or not history.get("epoch"):
        return
    df  = pd.DataFrame(history).dropna()
    fig = make_subplots(rows=1, cols=2, subplot_titles=("Accuracy", "Loss"),
                        horizontal_spacing=0.12)
    fig.add_trace(go.Scatter(x=df["epoch"], y=df["accuracy"], mode="lines+markers",
                             name="Train", line=dict(color="#34D399", width=2.5),
                             marker=dict(size=5)), row=1, col=1)
    if "val_accuracy" in df and df["val_accuracy"].notna().any():
        fig.add_trace(go.Scatter(x=df["epoch"], y=df["val_accuracy"], mode="lines+markers",
                                 name="Val", line=dict(color="#F4A261", width=2.5, dash="dash"),
                                 marker=dict(size=5)), row=1, col=1)
    fig.add_trace(go.Scatter(x=df["epoch"], y=df["loss"], mode="lines+markers",
                             name="Train Loss", line=dict(color="#F87171", width=2.5),
                             marker=dict(size=5), showlegend=False), row=1, col=2)
    if "val_loss" in df and df["val_loss"].notna().any():
        fig.add_trace(go.Scatter(x=df["epoch"], y=df["val_loss"], mode="lines+markers",
                                 name="Val Loss", line=dict(color="#818CF8", width=2.5, dash="dash"),
                                 marker=dict(size=5), showlegend=False), row=1, col=2)
    fig.update_layout(
        height=340, template="plotly_dark", hovermode="x unified",
        paper_bgcolor="rgba(0,0,0,0)", plot_bgcolor="rgba(255,255,255,0.02)",
        legend=dict(orientation="h", yanchor="bottom", y=1.08, xanchor="right", x=1,
                    bgcolor="rgba(0,0,0,0)", font=dict(color="#94A3B8")),
        margin=dict(l=20, r=20, t=50, b=20),
        font=dict(color="#94A3B8"),
    )
    fig.update_xaxes(title_text="Epoch", gridcolor="rgba(255,255,255,0.05)", zeroline=False)
    fig.update_yaxes(gridcolor="rgba(255,255,255,0.05)", zeroline=False)
    fig.update_yaxes(range=[0, 1], row=1, col=1)
    st.plotly_chart(fig, width="stretch", key="learning_curves")


def _save_learning_curves(history: dict, model_path: str = "",
                           output_dir: Path | None = None) -> str:
    out_dir   = output_dir if output_dir else _next_training_output_dir()
    stem      = Path(model_path).stem if model_path else "run"
    json_path = out_dir / f"history_{stem}.json"
    json_path.write_text(json.dumps(history))

    html_path = out_dir / f"curves_{stem}.html"
    df = pd.DataFrame(history).dropna()
    fig = make_subplots(rows=1, cols=2, subplot_titles=("Accuracy", "Loss"),
                        horizontal_spacing=0.12)
    fig.add_trace(go.Scatter(x=df["epoch"], y=df["accuracy"], mode="lines+markers",
                             name="Train Acc", line=dict(color="#34D399", width=2.5)), row=1, col=1)
    if "val_accuracy" in df and df["val_accuracy"].notna().any():
        fig.add_trace(go.Scatter(x=df["epoch"], y=df["val_accuracy"], mode="lines+markers",
                                 name="Val Acc", line=dict(color="#F4A261", width=2.5, dash="dash")),
                      row=1, col=1)
    fig.add_trace(go.Scatter(x=df["epoch"], y=df["loss"], mode="lines+markers",
                             name="Train Loss", line=dict(color="#F87171", width=2.5)), row=1, col=2)
    if "val_loss" in df and df["val_loss"].notna().any():
        fig.add_trace(go.Scatter(x=df["epoch"], y=df["val_loss"], mode="lines+markers",
                                 name="Val Loss", line=dict(color="#818CF8", width=2.5, dash="dash")),
                      row=1, col=2)
    fig.update_layout(height=420, template="plotly_white", hovermode="x unified",
                      margin=dict(l=30, r=20, t=60, b=30))
    fig.write_html(str(html_path))
    return str(html_path)


# ---------------------------------------------------------------------------
# Class balance
# ---------------------------------------------------------------------------

def _render_class_balance(balance: dict) -> None:
    ratio   = balance["imbalance_ratio"]
    classes = balance["classes"]
    counts  = balance["counts"]
    display = [DISEASE_DISPLAY.get(c, c) for c in classes]
    min_c   = min(counts.values())
    max_c   = max(counts.values())

    if balance["is_imbalanced"]:
        st.warning(f"Imbalance ratio {ratio:.1f}× — class weights will compensate during training.",
                   icon="⚠️")
    else:
        st.success(f"Classes are balanced (ratio {ratio:.1f}×).", icon="✅")

    mc1, mc2, mc3 = st.columns(3)
    mc1.metric("Total images",    balance["total"])
    mc2.metric("Classes",         len(classes))
    mc3.metric("Imbalance ratio", f"{ratio:.1f}×")

    colors = [
        "#E53E3E" if counts[c] == min_c else
        "#F4A261" if counts[c] == max_c else "#90CDF4"
        for c in classes
    ]
    fig = go.Figure(go.Bar(
        x=display, y=[counts[c] for c in classes],
        marker_color=colors,
        text=[counts[c] for c in classes], textposition="outside",
        hovertemplate="%{x}<br>%{y} images<extra></extra>",
    ))
    mean_count = balance["total"] / max(len(classes), 1)
    fig.add_hline(y=mean_count, line_dash="dash", line_color="#718096",
                  annotation_text=f"Mean ({mean_count:.0f})", annotation_position="top right")
    fig.update_layout(height=280, template="plotly_white",
                      margin=dict(l=20, r=20, t=20, b=60),
                      paper_bgcolor="rgba(0,0,0,0)",
                      xaxis=dict(tickangle=-30),
                      yaxis=dict(title="Images"))
    st.plotly_chart(fig, width="stretch", key="class_balance_chart")
    with st.expander("Class weights (used during training)"):
        wdf = pd.DataFrame([
            {"Class": DISEASE_DISPLAY.get(c, c), "Weight": f"{w:.3f}"}
            for c, w in balance["weights"].items()
        ])
        st.dataframe(wdf, hide_index=True, width="stretch")
        st.caption("Weight > 1 = under-represented class; model penalised more for missing it.")


# ---------------------------------------------------------------------------
# t-SNE visualisation
# ---------------------------------------------------------------------------

def _render_tsne(result: dict) -> None:
    class_names = result["class_names"]
    x, y, labels = result["x"], result["y"], result["labels"]

    fig = go.Figure()
    for i, cls in enumerate(class_names):
        mask = [j for j, lbl in enumerate(labels) if lbl == i]
        if not mask:
            continue
        fig.add_trace(go.Scatter(
            x=[x[j] for j in mask], y=[y[j] for j in mask],
            mode="markers",
            name=DISEASE_DISPLAY.get(cls, cls),
            marker=dict(color=_PALETTE[i % len(_PALETTE)], size=5,
                        opacity=0.75, line=dict(width=0)),
            hovertemplate=f"{DISEASE_DISPLAY.get(cls, cls)}<extra></extra>",
        ))
    fig.update_layout(
        height=480, template="plotly_white",
        legend=dict(orientation="v", x=1.01, y=1),
        margin=dict(l=20, r=160, t=20, b=20),
        paper_bgcolor="rgba(0,0,0,0)",
        xaxis=dict(showticklabels=False, zeroline=False, showgrid=False, title="t-SNE dim 1"),
        yaxis=dict(showticklabels=False, zeroline=False, showgrid=False, title="t-SNE dim 2"),
    )
    st.plotly_chart(fig, width="stretch", key="tsne_chart")
    st.caption(
        f"{result['n_samples']} samples plotted (of {result['n_total']} total). "
        "Tight, well-separated clusters = the backbone has learnt discriminative features."
    )


# ---------------------------------------------------------------------------
# Model comparison across training runs
# ---------------------------------------------------------------------------

def _build_excel(rows: list) -> bytes:
    """Build an Excel workbook from all training run histories and return as bytes."""
    import io
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()

    # ── Sheet 1: Summary ──────────────────────────────────────────────────
    ws = wb.active
    ws.title = "Summary"
    hdr_fill = PatternFill("solid", fgColor="C1121F")
    hdr_font = Font(bold=True, color="FFFFFF")

    summary_cols = ["Run", "Stage", "Epochs", "Best Epoch", "Best Val Acc", "Final Val Acc"]
    for col, name in enumerate(summary_cols, 1):
        cell = ws.cell(row=1, column=col, value=name)
        cell.fill = hdr_fill
        cell.font = hdr_font
        cell.alignment = Alignment(horizontal="center")

    for r_idx, row in enumerate(rows, 2):
        ws.cell(r_idx, 1, row["Run"])
        ws.cell(r_idx, 2, row["Stage"])
        ws.cell(r_idx, 3, row["Epochs"])
        ws.cell(r_idx, 4, row["Best Epoch"])
        ws.cell(r_idx, 5, row["Best Val Acc"])
        ws.cell(r_idx, 6, row["Final Val Acc"])

    for col in range(1, len(summary_cols) + 1):
        ws.column_dimensions[get_column_letter(col)].auto_size = True

    # ── Per-run sheets: one sheet per history file ────────────────────────
    epoch_cols = ["Epoch", "Train Accuracy", "Val Accuracy", "Train Loss", "Val Loss"]
    for row in rows:
        try:
            hist = json.loads(Path(row["_hist"]).read_text())
        except Exception:
            continue

        sheet_name = f"{row['Run']} {row['Stage']}"[:31]  # Excel limit
        ws2 = wb.create_sheet(title=sheet_name)

        for col, name in enumerate(epoch_cols, 1):
            cell = ws2.cell(row=1, column=col, value=name)
            cell.fill = hdr_fill
            cell.font = hdr_font
            cell.alignment = Alignment(horizontal="center")

        epochs    = hist.get("epoch", [])
        train_acc = hist.get("accuracy", [])
        val_acc   = hist.get("val_accuracy", [])
        train_los = hist.get("loss", [])
        val_los   = hist.get("val_loss", [])

        for i, ep in enumerate(epochs):
            ws2.cell(i + 2, 1, ep)
            ws2.cell(i + 2, 2, train_acc[i] if i < len(train_acc) else None)
            ws2.cell(i + 2, 3, val_acc[i]   if i < len(val_acc)   else None)
            ws2.cell(i + 2, 4, train_los[i] if i < len(train_los) else None)
            ws2.cell(i + 2, 5, val_los[i]   if i < len(val_los)   else None)

        for col in range(1, len(epoch_cols) + 1):
            ws2.column_dimensions[get_column_letter(col)].width = 18

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _render_model_comparison() -> None:
    outputs_base = Path("outputs")
    rows = []
    for run_dir in sorted(outputs_base.glob("Training*"), key=lambda p: p.name):
        for hist_file in sorted(run_dir.glob("history_*.json")):
            try:
                hist     = json.loads(hist_file.read_text())
                val_acc  = [v for v in hist.get("val_accuracy", []) if v is not None]
                val_loss = [v for v in hist.get("val_loss", []) if v is not None]
                epochs   = hist.get("epoch", [])
                best_val = max(val_acc) if val_acc else None
                best_ep  = epochs[val_acc.index(best_val)] if best_val and epochs else None
                rows.append({
                    "Run":           run_dir.name,
                    "Stage":         hist_file.stem.replace("history_", "").replace("_", " ").title(),
                    "Epochs":        len(epochs),
                    "Best Epoch":    best_ep or "—",
                    "Best Val Acc":  f"{best_val:.2%}" if best_val else "—",
                    "Final Val Acc": f"{val_acc[-1]:.2%}" if val_acc else "—",
                    "_best_val":     best_val or 0,
                    "_hist":         str(hist_file),
                })
            except Exception:
                pass

    if not rows:
        st.info("No training runs found. Complete a training run first.", icon="📊")
        return

    rows.sort(key=lambda r: r["_best_val"], reverse=True)
    best = rows[0]
    rc1, rc2, rc3 = st.columns(3)
    rc1.metric("Total runs",   len(set(r["Run"] for r in rows)))
    rc2.metric("Best val acc", best["Best Val Acc"])
    rc3.metric("Best run",     f"{best['Run']} / {best['Stage']}")

    display_cols = ["Run", "Stage", "Epochs", "Best Epoch", "Best Val Acc", "Final Val Acc"]
    st.dataframe([{k: r[k] for k in display_cols} for r in rows],
                 hide_index=True, width="stretch")

    # ── Excel export ──────────────────────────────────────────────────────
    st.download_button(
        label="Download results as Excel",
        data=_build_excel(rows),
        file_name="madomasi_training_results.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        type="secondary",
    )

    with st.expander("Val accuracy comparison chart"):
        fig = go.Figure()
        for i, r in enumerate(rows[:8]):
            try:
                hist = json.loads(Path(r["_hist"]).read_text())
                eps  = hist.get("epoch", [])
                va   = hist.get("val_accuracy", [])
                if eps and va:
                    fig.add_trace(go.Scatter(
                        x=eps, y=va, mode="lines+markers",
                        name=f"{r['Run']} / {r['Stage']}",
                        line=dict(color=_PALETTE[i % len(_PALETTE)], width=2),
                        marker=dict(size=5),
                    ))
            except Exception:
                pass
        fig.update_layout(
            height=320, template="plotly_white", hovermode="x unified",
            legend=dict(orientation="h", yanchor="bottom", y=1.02, xanchor="right", x=1),
            margin=dict(l=20, r=20, t=40, b=20),
            yaxis=dict(tickformat=".0%", range=[0, 1], title="Val Accuracy"),
            xaxis=dict(title="Epoch"),
        )
        st.plotly_chart(fig, width="stretch", key="model_comp_chart")

    keras_files = (
        sorted(Path("models/trained").rglob("*.keras"), key=lambda p: p.stat().st_mtime, reverse=True)
        if Path("models/trained").exists() else []
    )
    if keras_files:
        st.divider()
        st.caption("**Promote a model** — set it as the active model for inference and evaluation")
        sel = st.selectbox(
            "Model file", [str(p) for p in keras_files],
            format_func=lambda p: f"{Path(p).name}  ({Path(p).stat().st_size / 1e6:.1f} MB)",
            key="model_comp_select",
        )
        if st.button("Use this model", key="model_comp_promote", type="primary"):
            st.session_state["eval_model"]  = sel
            st.session_state["model_path"]  = sel
            st.session_state["tfl_model"]   = sel
            st.success(f"Active model set to `{Path(sel).name}`")


# ---------------------------------------------------------------------------
# Evaluation helpers
# ---------------------------------------------------------------------------

def _render_confusion_matrix(ev: dict) -> None:
    import numpy as _np

    cm          = _np.array(ev["confusion_matrix"])
    raw_names   = ev["class_names"]
    disp_names  = [DISEASE_DISPLAY.get(n, n) for n in raw_names]
    report      = ev.get("report", {})

    st.markdown("#### Confusion Matrix")

    row_sums = cm.sum(axis=1, keepdims=True).clip(min=1)
    cm_norm  = (cm / row_sums * 100).round(1)

    hover = [[
        f"Actual: {disp_names[r]}<br>Predicted: {disp_names[c]}<br>"
        f"{cm[r, c]} image(s) ({cm_norm[r, c]:.1f}%)"
        for c in range(len(disp_names))]
        for r in range(len(disp_names))
    ]

    fig = go.Figure(go.Heatmap(
        z=cm_norm, x=disp_names, y=disp_names,
        text=cm, texttemplate="%{text}",
        hoverinfo="text", hovertext=hover,
        colorscale="Blues", showscale=True,
        colorbar=dict(title="Recall %", ticksuffix="%"),
        zmin=0, zmax=100,
    ))
    fig.update_layout(
        height=520, template="plotly_white",
        xaxis=dict(title="Predicted", tickangle=-35, side="bottom"),
        yaxis=dict(title="Actual", autorange="reversed"),
        margin=dict(l=20, r=20, t=20, b=80),
        paper_bgcolor="rgba(0,0,0,0)",
    )
    st.plotly_chart(fig, width="stretch", key="cm_heatmap")
    st.caption("Cell values = raw count. Colour intensity = recall % (row-normalised).")

    if report:
        rows = []
        for name, dname in zip(raw_names, disp_names):
            r = report.get(name, {})
            rows.append({
                "Disease":   dname,
                "Precision": f"{r.get('precision', 0):.1%}",
                "Recall":    f"{r.get('recall', 0):.1%}",
                "F1":        f"{r.get('f1-score', 0):.1%}",
                "Support":   int(r.get("support", 0)),
            })
        with st.expander("Per-class metrics (Precision / Recall / F1)"):
            st.dataframe(rows, hide_index=True, width="stretch")


# ---------------------------------------------------------------------------
# Live training progress fragment (polls session state every 2 s)
# ---------------------------------------------------------------------------

@st.fragment(run_every="2s")
def _training_live_panel() -> None:
    """Renders live training/extraction progress without blocking the page.

    Checks _stage_state (extract/split) and _train_state (head/finetune) so
    every long-running stage keeps the panel alive while the server thread is free.
    """
    stage_state = st.session_state.get("_stage_state", {})
    train_state = st.session_state.get("_train_state", {})

    # ── Non-training stage (extract / split) ──────────────────────────────
    if stage_state.get("active"):
        st.progress(min(stage_state.get("progress", 0.0), 1.0),
                    text=stage_state.get("text", "Working…"))
        return

    # Show result once a non-training stage finishes
    stage_result = st.session_state.pop("_stage_result", None)
    if stage_result:
        if stage_result["success"]:
            st.success(stage_result["message"])
        else:
            st.error(f"Stage failed: {stage_result['message']}")

    # ── Keras training stage (head / finetune) ────────────────────────────
    if not train_state or not train_state.get("history", {}).get("epoch"):
        from components.ui_helpers import load_latest_history
        hist = load_latest_history() or st.session_state.get("train_history")
        if hist and hist.get("epoch"):
            plot_learning_curves(hist)
        else:
            html(
                '<div class="empty-state">'
                '<span class="empty-state-icon">📈</span>'
                '<div class="empty-state-title">No training history yet</div>'
                '<div class="empty-state-sub">Run a training stage — live curves appear here</div>'
                '</div>'
            )
        return

    progress = train_state.get("progress", 0.0)
    text     = train_state.get("text", "")
    history  = train_state.get("history", {})
    logs     = train_state.get("logs", {})
    active   = train_state.get("active", False)

    st.progress(min(progress, 1.0), text=text)

    if logs:
        acc      = logs.get("accuracy")
        loss_val = logs.get("loss")
        val_acc  = logs.get("val_accuracy")
        val_loss = logs.get("val_loss")
        m1, m2, m3, m4 = st.columns(4)
        m1.metric("Train Acc",  f"{acc:.3f}"      if acc      is not None else "—")
        m2.metric("Train Loss", f"{loss_val:.4f}" if loss_val is not None else "—")
        m3.metric("Val Acc",    f"{val_acc:.3f}"  if val_acc  is not None else "—")
        m4.metric("Val Loss",   f"{val_loss:.4f}" if val_loss is not None else "—")

    if history.get("epoch"):
        plot_learning_curves(history)

    if not active:
        result = st.session_state.pop("_train_result", None)
        if result:
            if result["success"]:
                st.success(result["message"])
            else:
                st.error(f"Stage failed: {result['message']}")


# ---------------------------------------------------------------------------
# Tab entry point
# ---------------------------------------------------------------------------

def render() -> None:
    """Render the full Training tab."""
    from components.system_dashboard import record_resource_sample

    st.session_state.pop("_models_dirty", False)

    # ── Session defaults ──────────────────────────────────────────────────────
    if "shared_data_dir" not in st.session_state:
        st.session_state["shared_data_dir"] = detect_dir([
            st.session_state.get("split_train_dir", ""),
            "data/splits/train", "data/raw/raw/tomato", "data/raw",
        ])
    if "shared_model_base" not in st.session_state:
        st.session_state["shared_model_base"] = detect_dir(["models/trained"], fallback="")
    if "features_dir" not in st.session_state:
        detected = detect_dir(["data/features"], fallback="")
        if detected:
            st.session_state["features_dir"] = detected
    if "head_full_path" not in st.session_state:
        for c in ["models/trained/head_full.keras", "models/trained/latest.keras"]:
            if Path(c).exists():
                st.session_state["head_full_path"] = c
                break

    # ── Shared config bar ─────────────────────────────────────────────────────
    cfg1, cfg2, cfg3 = st.columns([2, 2, 1])
    data_dir        = cfg1.text_input("Dataset directory", placeholder="data/splits/train",
                                      key="shared_data_dir")
    model_save_base = cfg2.text_input("Model output", placeholder="models/trained",
                                      key="shared_model_base")
    batch_size      = cfg3.select_slider("Batch size", options=[4, 8, 16, 32, 64], value=16)

    col_left, col_right = st.columns([2, 3], gap="large")

    # ── Right: live training output ───────────────────────────────────────────
    with col_right:
        st.subheader("Live Training Curves")
        _training_live_panel()

    # ── Left: pipeline stages ─────────────────────────────────────────────────
    with col_left:
        split_done = bool(st.session_state.get("split_train_dir"))
        feat_done  = bool(st.session_state.get("features_dir"))
        head_done  = bool(st.session_state.get("head_full_path"))
        ft_done    = bool(st.session_state.get("fine_tuned_path"))

        # Stage 0 — Split
        with st.expander(f"{'✅' if split_done else '○'} Stage 0 — Split Dataset",
                         expanded=not split_done):
            if "raw_source" not in st.session_state:
                st.session_state["raw_source"] = detect_dir([
                    "data/raw/raw/tomato", "data/raw/tomato", "data/raw", "dataset",
                ])
            if "split_output" not in st.session_state:
                st.session_state["split_output"] = detect_dir(["data/splits"], fallback="")
            raw_source   = st.text_input("Raw dataset dir", placeholder="path/to/raw", key="raw_source")
            split_output = st.text_input("Splits output dir", placeholder="data/splits", key="split_output")
            r1, r2, r3   = st.columns(3)
            train_ratio  = r1.number_input("Train %", 0, 100, 70, step=5, key="train_ratio") / 100
            val_ratio    = r2.number_input("Val %",   0, 100, 20, step=5, key="val_ratio")   / 100
            test_ratio   = r3.number_input("Test %",  0, 100, 10, step=5, key="test_ratio")  / 100
            ratio_ok = abs(train_ratio + val_ratio + test_ratio - 1.0) < 1e-6
            if not ratio_ok:
                st.warning("Train + Val + Test must sum to 100%.")
            btn_split = st.button(
                "Split Dataset", type="primary", width="stretch",
                disabled=not ratio_ok or st.session_state.get("training_active", False),
                key="btn_split",
            )
            if split_done:
                st.caption(f"Train: `{st.session_state.split_train_dir}`")

        # Stage 1 — Feature Extraction
        with st.expander(f"{'✅' if feat_done else '○'} Stage 1 — Feature Extraction",
                         expanded=split_done and not feat_done):
            _src_opts  = list(dict.fromkeys([
                p for p in [
                    st.session_state.get("split_train_dir", ""),
                    "data/splits/train", "data/raw/raw/tomato", data_dir,
                ] if p and Path(p).exists()
            ]))
            _CUSTOM_SRC = "Custom path..."
            if _src_opts:
                _choice = st.selectbox(
                    "Source data", options=_src_opts + [_CUSTOM_SRC],
                    format_func=lambda x: x if x == _CUSTOM_SRC else Path(x).as_posix(),
                    key="feat_src_select",
                )
                feat_src_dir = (
                    st.text_input("Custom source dir", placeholder="data/splits/train", key="feat_src_custom")
                    if _choice == _CUSTOM_SRC else _choice
                )
            else:
                feat_src_dir = st.text_input("Source data dir", placeholder="data/splits/train",
                                             key="feat_src_custom")

            if "feat_out" not in st.session_state:
                st.session_state["feat_out"] = detect_dir(["data/features"], fallback="")
            feat_out   = st.text_input("Features output dir", placeholder="data/features", key="feat_out")
            feat_batch = st.select_slider("Batch size", options=[4, 8, 16, 32, 64],
                                          value=batch_size, key="feat_batch_size")
            btn_extract = st.button(
                "Extract Features", type="primary", width="stretch",
                disabled=st.session_state.get("training_active", False), key="btn_extract",
            )
            if feat_done:
                st.caption(f"Features: `{st.session_state.features_dir}`")
                if st.button("Visualise Feature Space (t-SNE)", key="btn_tsne", type="secondary"):
                    with st.spinner("Running t-SNE — this takes ~30 s for 1 500 samples…"):
                        try:
                            st.session_state.tsne_result = compute_tsne(st.session_state.features_dir)
                        except Exception as _e:
                            st.error(f"t-SNE failed: {_e}")
                if st.session_state.get("tsne_result"):
                    _render_tsne(st.session_state.tsne_result)

            if st.session_state.get("class_balance") and not feat_done:
                st.divider()
                st.caption("**Training data balance check**")
                _render_class_balance(st.session_state.class_balance)

        # Stage 2 — Train Head
        with st.expander(f"{'✅' if head_done else '○'} Stage 2 — Train Head",
                         expanded=feat_done and not head_done):
            head_epochs = st.slider("Epochs", 1, 50, 10, key="head_epochs")
            head_lr     = st.select_slider("Learning rate", [0.0001, 0.001, 0.01, 0.1],
                                           value=0.001, key="head_lr")
            btn_head = st.button(
                "Train Head", type="primary", width="stretch",
                disabled=not feat_done or st.session_state.get("training_active", False),
                key="btn_head",
            )
            if not feat_done:
                st.caption("Complete Feature Extraction first.")
            elif head_done:
                st.caption(f"Model: `{st.session_state.head_full_path}`")

        # Stage 3 — Fine-Tuning
        with st.expander(f"{'✅' if ft_done else '○'} Stage 3 — Fine-Tuning",
                         expanded=head_done and not ft_done):
            ft_epochs = st.slider("Epochs", 1, 50, 10, key="ft_epochs")
            ft_lr     = st.select_slider("Learning rate", [0.000001, 0.00001, 0.0001],
                                         value=0.00001, key="ft_lr")
            btn_ft = st.button(
                "Fine-Tune", type="primary", width="stretch",
                disabled=not head_done or st.session_state.get("training_active", False),
                key="btn_ft",
            )
            if not head_done:
                st.caption("Train the Head first.")
            elif ft_done:
                st.caption(f"Model: `{st.session_state.fine_tuned_path}`")

    # ── Stage execution (runs after button clicks, outside column context) ─────
    if btn_split:
        if not raw_source:
            st.error("Enter the raw dataset directory before splitting.")
        elif not Path(raw_source).exists():
            st.error(f"Directory not found: `{raw_source}`")
        else:
            st.session_state.update({
                "training_active": True, "_active_stage": "split",
                "_raw_source": raw_source,
                "_split_output": split_output or "data/splits",
                "_train_ratio": train_ratio, "_val_ratio": val_ratio, "_test_ratio": test_ratio,
            })
            st.rerun()

    if btn_extract:
        _ext_data = feat_src_dir or st.session_state.get("split_train_dir") or data_dir
        if not _ext_data:
            st.error("Select a source data directory.")
        elif not Path(_ext_data).exists():
            st.error(f"Directory not found: `{_ext_data}`")
        else:
            st.session_state.update({
                "training_active": True, "_active_stage": "extract",
                "_data_dir": _ext_data,
                "_feat_out": feat_out or "data/features",
                "_batch_size": feat_batch,
            })
            st.rerun()

    if btn_head:
        _feat_dir = st.session_state.get("features_dir")
        if not _feat_dir or not Path(_feat_dir).exists():
            st.error("Features directory not found. Run Feature Extraction first.")
        elif "_class_weights" not in st.session_state:
            st.warning("Run the data split (Stage 0) first to compute class weights.")
            st.stop()
        else:
            st.session_state.update({
                "training_active": True, "_active_stage": "head",
                "_head_epochs": head_epochs, "_head_lr": head_lr,
                "_model_save_base": model_save_base or "models/trained/latest.keras",
            })
            st.rerun()

    if btn_ft:
        _head_path = st.session_state.get("head_full_path")
        _ft_data   = data_dir or st.session_state.get("split_train_dir", "")
        if not _head_path or not Path(_head_path).exists():
            st.error("Head model not found. Complete Stage 2 first.")
        elif not _ft_data or not Path(_ft_data).exists():
            st.error("Enter a valid dataset directory for fine-tuning.")
        else:
            st.session_state.update({
                "training_active": True, "_active_stage": "finetune",
                "_ft_epochs": ft_epochs, "_ft_lr": ft_lr,
                "_batch_size": batch_size, "_data_dir": _ft_data,
            })
            st.rerun()

    # ── Active training ────────────────────────────────────────────────────────
    if st.session_state.get("training_active", False):
        stage = st.session_state.get("_active_stage", "")
        record_resource_sample("Training")

        from streamlit.runtime.scriptrunner import add_script_run_ctx, get_script_run_ctx
        ctx = get_script_run_ctx()

        if stage == "split" and not st.session_state.get("_train_thread_active"):
            if _TRAIN_LOCK.acquire(blocking=False):
                st.session_state["_train_thread_active"] = True
                st.session_state["_stage_state"] = {"progress": 0.0, "text": "Initialising split…", "active": True}

                def _run_split():
                    _success = False
                    def _split_prog(current, total, class_name):
                        st.session_state["_stage_state"] = {
                            "progress": current / max(total, 1),
                            "text":     f"Splitting {current}/{total}: {class_name}",
                            "active":   True,
                        }
                    try:
                        result = split_dataset(
                            source_dir  = st.session_state["_raw_source"],
                            output_dir  = st.session_state["_split_output"],
                            train_ratio = st.session_state["_train_ratio"],
                            val_ratio   = st.session_state["_val_ratio"],
                            test_ratio  = st.session_state["_test_ratio"],
                            progress_fn = _split_prog,
                        )
                        st.session_state.split_train_dir = result["train_dir"]
                        st.session_state.split_val_dir   = result["val_dir"]
                        st.session_state.split_test_dir  = result["test_dir"]
                        st.session_state.pop("shared_data_dir", None)
                        try:
                            balance     = check_class_balance(result["train_dir"])
                            st.session_state.class_balance = balance
                            class_names = sorted(balance["classes"])
                            st.session_state["_class_weights"] = {
                                i: balance["weights"][cls] for i, cls in enumerate(class_names)
                            }
                        except Exception:
                            pass
                        c = result["counts"]
                        st.session_state["_stage_result"] = {
                            "success": True,
                            "message": (
                                f"{result['num_classes']} classes — "
                                f"train: {c['train']} · val: {c['val']} · test: {c['test']} images"
                            ),
                        }
                        _success = True
                    except Exception as exc:
                        st.session_state["_stage_result"] = {"success": False, "message": str(exc)}
                    finally:
                        _TRAIN_LOCK.release()
                        load_cached_model.clear()
                        st.session_state["_models_dirty"]        = True
                        st.session_state["training_active"]      = False
                        st.session_state["_train_thread_active"] = False
                        st.session_state["_stage_state"] = {
                            "progress": 1.0,
                            "text":     "Done" if _success else "Failed",
                            "active":   False,
                        }

                t = threading.Thread(target=_run_split, daemon=True)
                add_script_run_ctx(t, ctx)
                t.start()

        elif stage == "extract" and not st.session_state.get("_train_thread_active"):
            if _TRAIN_LOCK.acquire(blocking=False):
                st.session_state["_train_thread_active"] = True
                st.session_state["_stage_state"] = {"progress": 0.0, "text": "Initialising extraction…", "active": True}

                def _run_extract():
                    _success = False
                    def _ext_prog(current, total):
                        st.session_state["_stage_state"] = {
                            "progress": min(current / max(total, 1), 1.0),
                            "text":     f"Extracting batch {current}/{total}",
                            "active":   True,
                        }
                    try:
                        result = extract_features(
                            data_dir    = st.session_state["_data_dir"],
                            output_dir  = st.session_state["_feat_out"],
                            batch_size  = st.session_state["_batch_size"],
                            progress_fn = _ext_prog,
                        )
                        st.session_state.features_dir = result["features_dir"]
                        st.session_state["_stage_result"] = {
                            "success": True,
                            "message": (
                                f"{result['num_samples']} samples · {result['num_classes']} classes · "
                                f"shape {result['feature_shape']} → `{result['features_dir']}`"
                            ),
                        }
                        _success = True
                    except Exception as exc:
                        st.session_state["_stage_result"] = {"success": False, "message": str(exc)}
                    finally:
                        _TRAIN_LOCK.release()
                        load_cached_model.clear()
                        st.session_state["_models_dirty"]        = True
                        st.session_state["training_active"]      = False
                        st.session_state["_train_thread_active"] = False
                        st.session_state["_stage_state"] = {
                            "progress": 1.0,
                            "text":     "Done" if _success else "Failed",
                            "active":   False,
                        }

                t = threading.Thread(target=_run_extract, daemon=True)
                add_script_run_ctx(t, ctx)
                t.start()

        elif stage in ("head", "finetune"):
            # Run model.fit() in a background thread so the Streamlit server thread
            # stays free — this keeps the System Dashboard responsive during training.
            if not st.session_state.get("_train_thread_active"):
                if _TRAIN_LOCK.acquire(blocking=False):
                    st.session_state["_train_thread_active"] = True

                    if stage == "head":
                        _cw = st.session_state.get("_class_weights")

                        def _run():
                            try:
                                cb = StreamlitTrainCallback(stage_label="Train Head")
                                result = train_head(
                                    features_dir  = st.session_state["features_dir"],
                                    epochs        = st.session_state["_head_epochs"],
                                    lr            = st.session_state["_head_lr"],
                                    output_path   = st.session_state["_model_save_base"],
                                    callbacks     = [cb],
                                    class_weights = _cw,
                                )
                                st.session_state.head_full_path = result["full_model_path"]
                                st.session_state.train_history  = cb.history
                                record_resource_sample("Training")
                                _save_learning_curves(cb.history, result["full_model_path"],
                                                      output_dir=cb.output_dir)
                                st.session_state["_train_result"] = {
                                    "success": True,
                                    "message": f"Head trained → `{result['full_model_path']}`",
                                }
                            except Exception as exc:
                                st.session_state["_train_result"] = {"success": False, "message": str(exc)}
                            finally:
                                _TRAIN_LOCK.release()
                                load_cached_model.clear()
                                st.session_state["_models_dirty"]        = True
                                st.session_state["training_active"]      = False
                                st.session_state["_train_thread_active"] = False
                                state = st.session_state.get("_train_state", {})
                                st.session_state["_train_state"] = {**state, "active": False}
                    else:
                        _cw = st.session_state.get("_class_weights")

                        def _run():
                            try:
                                cb = StreamlitTrainCallback(stage_label="Fine-Tuning")
                                result = fine_tune_model(
                                    model_path    = st.session_state["head_full_path"],
                                    data_dir      = st.session_state["_data_dir"],
                                    batch_size    = st.session_state["_batch_size"],
                                    epochs        = st.session_state["_ft_epochs"],
                                    lr            = st.session_state["_ft_lr"],
                                    callbacks     = [cb],
                                    class_weights = _cw,
                                )
                                st.session_state.fine_tuned_path = result["path"]
                                st.session_state.train_history   = cb.history
                                record_resource_sample("Training")
                                _save_learning_curves(cb.history, result["path"], output_dir=cb.output_dir)
                                st.session_state["_train_result"] = {
                                    "success": True,
                                    "message": f"Fine-tuned model → `{result['path']}`",
                                }
                            except Exception as exc:
                                st.session_state["_train_result"] = {"success": False, "message": str(exc)}
                            finally:
                                _TRAIN_LOCK.release()
                                load_cached_model.clear()
                                st.session_state["_models_dirty"]        = True
                                st.session_state["training_active"]      = False
                                st.session_state["_train_thread_active"] = False
                                state = st.session_state.get("_train_state", {})
                                st.session_state["_train_state"] = {**state, "active": False}

                    t = threading.Thread(target=_run, daemon=True)
                    add_script_run_ctx(t, ctx)
                    t.start()

    # ── Evaluation ────────────────────────────────────────────────────────────
    eval_done = bool(st.session_state.get("eval_result"))
    with st.expander(f"{'✅' if eval_done else '○'} Model Evaluation", expanded=False):
        ev_c1, ev_c2 = st.columns(2)
        with ev_c1:
            eval_model_path = model_picker("Model to evaluate", key="eval_model")
            if "eval_test_dir" not in st.session_state:
                st.session_state["eval_test_dir"] = detect_dir([
                    st.session_state.get("split_test_dir", ""),
                    "data/splits/test", "data/splits/val",
                ])
            eval_test_dir = st.text_input("Test directory", placeholder="data/splits/test",
                                          key="eval_test_dir")
            eval_btn = st.button("Evaluate", type="primary", width="stretch")

        with ev_c2:
            if eval_btn:
                if not Path(eval_model_path).exists():
                    st.error("Model file not found.")
                elif not Path(eval_test_dir).exists():
                    st.error("Test directory not found.")
                else:
                    with st.spinner("Evaluating…"):
                        try:
                            ev = evaluate_model(eval_model_path, eval_test_dir)
                            st.session_state.eval_result     = ev
                            st.session_state.eval_model_name = Path(eval_model_path).name
                        except Exception as e:
                            st.error(f"Evaluation failed: {e}")

            if st.session_state.get("eval_result"):
                ev  = st.session_state.eval_result
                acc = ev["accuracy"]
                m1, m2 = st.columns(2)
                m1.metric("Test Accuracy", f"{acc:.2%}")
                m2.metric("Test Loss",     f"{ev['loss']:.4f}")
                verdict = (
                    "Excellent — ready for deployment." if acc >= 0.90 else
                    "Good — consider more fine-tuning."  if acc >= 0.75 else
                    "Needs improvement — increase epochs or data."
                )
                st.caption(verdict)

        if st.session_state.get("eval_result"):
            ev = st.session_state.eval_result
            if "confusion_matrix" in ev:
                _render_confusion_matrix(ev)

    # ── TFLite Export ─────────────────────────────────────────────────────────
    tfl_done = bool(st.session_state.get("tfl_exported"))
    with st.expander(f"{'✅' if tfl_done else '○'} Export for Raspberry Pi (TFLite)",
                     expanded=False):
        tfl_c1, tfl_c2 = st.columns(2)
        with tfl_c1:
            tfl_model = model_picker("Model to export", key="tfl_model")
            tfl_out   = st.text_input("Output path", value="models/trained/model.tflite",
                                      key="tfl_out_path")
            tfl_btn   = st.button("Export TFLite (INT8)", type="primary", width="stretch")
        with tfl_c2:
            if tfl_btn:
                if not bool(tfl_model) or not Path(tfl_model).exists():
                    st.error("Model file not found.")
                else:
                    with st.spinner("Converting to TFLite with INT8 quantisation…"):
                        try:
                            out     = convert_model(tfl_model, tfl_out)
                            size_mb = Path(out).stat().st_size / 1e6
                            st.session_state.tfl_exported = True
                            st.success(f"Exported → `{out}` ({size_mb:.1f} MB)")
                        except Exception as e:
                            st.error(f"Export failed: {e}")

    # ── Model Comparison ──────────────────────────────────────────────────────
    with st.expander("📊 Training Runs — Model Comparison", expanded=False):
        _render_model_comparison()
